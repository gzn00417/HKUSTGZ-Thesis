
# ═══════════════════════════════════════════════════════════════════════════════
# DESIGN SYSTEM · REFERENCE LIBRARY
# build.py — HKUST(GZ) PhD Thesis Defence Deck
#
# PURPOSE   Authoritative design token registry and component library.
#           Copy functions into a new deck builder, or import as a module.
#           Never add slide content here — keep it functions-only.
# WORKFLOW  python3 build.py  →  empty PPTX to smoke-test the design system.
#           Actual deck builder (v0.4) was a one-shot script; v0.5+ should
#           be built from this design system as a separate build script.
#
# LAYERS
#   1  Primitives   _bg · rect · soft_shadow · hline · txt · rich · bullets · pic · notes
#   2  Cards        info_card · result_card · warn_card · gold_card · gray_card · stat_card
#   3  Widgets      chip · arrow · mathbox
#   4  Structural   progbar · footer · base · take · caption · flatcol · add_table · section
#   5  Appendix     bslide  (B-numbered backup slides)
#   6  Utilities    _force · _tf · _walk · _register_notes_master  (font + package fix)
# ═══════════════════════════════════════════════════════════════════════════════

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

# ── Palette ───────────────────────────────────────────────────────────────────
# Primary: NAVY (headings), GOLD (accents/underlines), DARK (body text), MUTED (secondary)
# Semantic: RED (warn), GREEN (result), AMBER (gold-card label)
# Light variants (_L) = card backgrounds; Border variants (_BD) = card borders
NAVY   = RGBColor(0x00, 0x33, 0x65)
NAVY2  = RGBColor(0x13, 0x2F, 0x57)
GOLD   = RGBColor(0xD4, 0xAF, 0x37)
GOLD2  = RGBColor(0xC0, 0x92, 0x1C)
RED    = RGBColor(0xB2, 0x2A, 0x2A)
GREEN  = RGBColor(0x05, 0x7A, 0x55)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x1C, 0x27, 0x39)
MUTED  = RGBColor(0x55, 0x66, 0x7D)
CREAM  = RGBColor(0xC3, 0xD3, 0xE6)

NAVY_L  = RGBColor(0xEC, 0xF3, 0xFC)
GREEN_L = RGBColor(0xEA, 0xF8, 0xF2)
RED_L   = RGBColor(0xFB, 0xEF, 0xEF)
GOLD_L  = RGBColor(0xFD, 0xF7, 0xE7)
GRAY_L  = RGBColor(0xF0, 0xF3, 0xF8)
NAVY_BD = RGBColor(0xBE, 0xD6, 0xF1)
GREEN_BD= RGBColor(0x9A, 0xE0, 0xCA)
RED_BD  = RGBColor(0xE6, 0xB6, 0xB6)
GOLD_BD = RGBColor(0xE8, 0xD5, 0x95)
GRAY_BD = RGBColor(0xD2, 0xDB, 0xE6)
AMBER   = RGBColor(0x86, 0x5F, 0x00)
PROG_OFF= RGBColor(0x9A, 0xA6, 0xB5)
HAIR    = RGBColor(0xDD, 0xE3, 0xEC)

FONT = "Times New Roman"
FIG  = "figures"

# ── Type scale ────────────────────────────────────────────────────────────────
# 4 prose tiers + display tier + 1 table size. Do NOT invent new sizes.
# T4=12 footer/caption · T3=16 card/take/chip · T2=20 body/claim · T1=28 title
# Display (hero glyphs only): S_BIG=40 cover/metric · S_QA=54 Q&A · S_SECNUM=96 section
# Tables: S_TABLE=14 (single fixed size for data density)
S_FOOT  = 12   # T4
S_PROG  = 12   # T4
S_CAP   = 12   # T4
S_CARDB = 16   # T3
S_CARDL = 16   # T3
S_BODY  = 20   # T2
S_CLAIM = 20   # T2
S_TITLE = 28   # T1
S_KICK  = 28   # stat-card hero number → reuse T1
S_BIG   = 40   # display
S_FORM  = 16   # formulas → T3
S_TAKE  = 16   # T3
S_SECTL = 28   # T1 (section title)
S_SECNUM= 96   # display
S_COVER = 40   # display
S_QA    = 54   # display
S_TABLE = 14   # tables only

# ── Layout constants ──────────────────────────────────────────────────────────
# Slide: 13.33" × 7.50"  (16:9 widescreen)
# Content zone: x ∈ [0.55, 12.78]  y ∈ [1.98, 6.50]  (below title+claim, above take/footer)
# Two-column gap = 0.15" → col_w = 6.04"
# Three-column gap = 0.22" → col_w = 3.93"
CONTENT_X0    = 0.55
CONTENT_W     = 12.23
CONTENT_TOP   = 1.98
CONTENT_MAX_Y = 6.50
TAKE_Y_MAX    = 6.55
FOOT_RULE     = 7.16
COL2_GAP      = 0.15
COL3_GAP      = 0.22
CAP_PAD       = 0.06

# Navigation bar sections (order = left→right progbar order)
SECTIONS = ["Introduction", "Semantic Scaling", "Observational",
            "Search Scaling", "Conclusion"]

# ── Presentation setup ────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.50)
BLANK = prs.slide_layouts[6]   # blank layout — all content placed manually


# ═══════════════════════════════════════════════════════════════════════════════
# LAYER 1 · PRIMITIVES
# Low-level wrappers around python-pptx. All coordinates in inches, sizes in pt.
# ═══════════════════════════════════════════════════════════════════════════════

def _bg(s, c=WHITE):
    f = s.background.fill; f.solid(); f.fore_color.rgb = c


def rect(s, l, t, w, h, color, line_color=None, line_pt=0.75, rounded=False, adj=0.08):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    if rounded:
        try: shp.adjustments[0] = adj
        except Exception: pass
    shp.shadow.inherit = False
    if color is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = color
    if line_color: shp.line.color.rgb = line_color; shp.line.width = Pt(line_pt)
    else: shp.line.fill.background()
    return shp


def soft_shadow(shp, blur=0.030, dist=0.018, alpha=38, clr="9AA6B5"):
    spPr = shp._element.spPr
    for e in spPr.findall(qn("a:effectLst")): spPr.remove(e)
    el = spPr.makeelement(qn("a:effectLst"), {})
    sh = el.makeelement(qn("a:outerShdw"),
                        {"blurRad": str(int(blur*914400)), "dist": str(int(dist*914400)),
                         "dir": "5400000", "rotWithShape": "0"})
    c = sh.makeelement(qn("a:srgbClr"), {"val": clr})
    c.append(c.makeelement(qn("a:alpha"), {"val": str(alpha*1000)}))
    sh.append(c); el.append(sh); spPr.append(el)
    return shp


def hline(s, l, t, w, color=GOLD, pt=1.4):
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Pt(pt))
    shp.shadow.inherit = False
    shp.fill.solid(); shp.fill.fore_color.rgb = color; shp.line.fill.background()
    return shp


def txt(s, l, t, w, h, text, size=S_BODY, bold=False, color=DARK,
        align=PP_ALIGN.LEFT, italic=False, wrap=True, anchor=None, spacing=1.02):
    b = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = b.text_frame; tf.word_wrap = wrap
    tf.margin_left = Pt(2); tf.margin_right = Pt(2); tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    if anchor is not None: tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.text = text; p.alignment = align; p.line_spacing = spacing
    p.font.size = Pt(size); p.font.bold = bold; p.font.italic = italic
    p.font.color.rgb = color; p.font.name = FONT
    return b


def rich(s, l, t, w, h, segs, size=S_BODY, align=PP_ALIGN.LEFT, anchor=None, spacing=1.02):
    # segs: list of (text, bold, color) or (text, bold, color, italic)
    b = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = b.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    if anchor is not None: tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = spacing
    for sg in segs:
        text, bold, color = sg[0], sg[1], sg[2]
        italic = sg[3] if len(sg) > 3 else False
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = FONT
    return b


def bullets(s, l, t, w, h, items, size=S_BODY, color=DARK, spacing=8, sub=MUTED):
    # items: str  OR  (text, level)  OR  (text, level, bold)
    # level 0 → bullet "•", level 1 → sub-bullet "–" at size-2
    b = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = b.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    for i, it in enumerate(items):
        lvl, text, bd = 0, it, False
        if isinstance(it, tuple):
            if len(it) == 3: text, lvl, bd = it
            else: text, lvl = it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("•  " + text) if lvl == 0 else ("–  " + text)
        p.font.size = Pt(size if lvl == 0 else size - 2)
        p.font.bold = bd
        p.font.color.rgb = color if lvl == 0 else sub
        p.font.name = FONT; p.space_after = Pt(spacing if lvl == 0 else 3); p.line_spacing = 1.02
    return b


def pic(s, fname, l, t, width=None, height=None, caption=None):
    # fname: filename without extension, resolved from FIG directory
    # Returns the rendered height in inches (useful for positioning caption)
    path = os.path.join(FIG, fname + ".png")
    kw = {}
    if width:  kw["width"]  = Inches(width)
    if height: kw["height"] = Inches(height)
    if not kw: kw["width"]  = Inches(8)
    p = s.shapes.add_picture(path, Inches(l), Inches(t), **kw)
    rh = p.height / 914400.0
    rw = p.width  / 914400.0
    if caption:
        txt(s, l, t + rh + CAP_PAD, rw, 0.34, caption,
            S_CAP, False, MUTED, PP_ALIGN.CENTER, italic=True)
    return rh


def notes(s, text): s.notes_slide.notes_text_frame.text = text


# ═══════════════════════════════════════════════════════════════════════════════
# LAYER 2 · CARDS
# Rounded-rect cards with a left color strip, optional label + body.
# Five semantic variants — pick by meaning, not color preference.
#   info_card   → NAVY   (neutral information / context)
#   result_card → GREEN  (findings / positive outcomes)
#   warn_card   → RED    (limitations / caveats)
#   gold_card   → GOLD   (key takeaways / thesis claims)
#   gray_card   → MUTED  (secondary / appendix content)
# stat_card: hero number + label + sentence (for big-metric callouts)
# ═══════════════════════════════════════════════════════════════════════════════

def _card(s, l, t, w, h, label, body, bg, bd, strip, lc, body_col=DARK, body_size=S_CARDB):
    rect(s, l, t, w, h, bg, line_color=bd, line_pt=1.0, rounded=True)
    rect(s, l + 0.02, t + 0.11, 0.09, h - 0.22, strip)
    if label:
        txt(s, l + 0.26, t + 0.11, w - 0.38, 0.34, label, S_CARDL, True, lc)
    if body:
        top = t + 0.50 if label else t + 0.15
        ht  = h - 0.60 if label else h - 0.26
        if isinstance(body, list):
            bullets(s, l + 0.26, top, w - 0.46, ht, body, body_size, body_col, spacing=5)
        else:
            txt(s, l + 0.26, top, w - 0.46, ht, body, body_size, False, body_col)

def info_card(s,l,t,w,h,lab,body,**k):   _card(s,l,t,w,h,lab,body,NAVY_L,NAVY_BD,NAVY,NAVY,**k)
def result_card(s,l,t,w,h,lab,body,**k): _card(s,l,t,w,h,lab,body,GREEN_L,GREEN_BD,GREEN,GREEN,**k)
def warn_card(s,l,t,w,h,lab,body,**k):   _card(s,l,t,w,h,lab,body,RED_L,RED_BD,RED,RED,**k)
def gold_card(s,l,t,w,h,lab,body,**k):   _card(s,l,t,w,h,lab,body,GOLD_L,GOLD_BD,GOLD,AMBER,**k)
def gray_card(s,l,t,w,h,lab,body,**k):   _card(s,l,t,w,h,lab,body,GRAY_L,GRAY_BD,MUTED,MUTED,**k)


def stat_card(s, l, t, w, h, big, label, sentence, accent=NAVY, bg=NAVY_L, bd=NAVY_BD):
    # big: large hero string (e.g. "+3.2%")  label: short tag  sentence: 1-line explanation
    rect(s, l, t, w, h, bg, line_color=bd, line_pt=1.0, rounded=True)
    txt(s, l + 0.22, t + 0.16, w - 0.40, 0.66, big, S_KICK, True, accent)
    txt(s, l + 0.22, t + 0.86, w - 0.40, 0.36, label, S_CARDL, True, DARK)
    txt(s, l + 0.22, t + 1.26, w - 0.40, h - 1.36, sentence, S_CARDB, False, MUTED)


# ═══════════════════════════════════════════════════════════════════════════════
# LAYER 3 · WIDGETS
# Standalone inline elements: chips, arrows, formula boxes.
# ═══════════════════════════════════════════════════════════════════════════════

def chip(s, l, t, w, h, text, fill, tc, size=S_CARDB, bold=True, bd=None, italic=False):
    # Rounded pill label. Used for pipeline stage tags and progbar segments.
    r = rect(s, l, t, w, h, fill, line_color=bd, line_pt=1.0, rounded=True)
    tf = r.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(4); tf.margin_right = Pt(4); tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.text = text
    p.font.size = Pt(size); p.font.bold = bold; p.font.italic = italic
    p.font.color.rgb = tc; p.font.name = FONT
    return r


def arrow(s, l, t, text="→", size=20, color=GOLD2):
    txt(s, l, t, 0.44, 0.44, text, size, True, color, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def mathbox(s, l, t, w, h, expr, label=None, expr_size=S_FORM, accent=NAVY):
    # Gray rounded box for formulas/equations. label optional (pinned top-left).
    rect(s, l, t, w, h, GRAY_L, line_color=GRAY_BD, line_pt=1.0, rounded=True)
    rect(s, l + 0.02, t + 0.10, 0.09, h - 0.20, accent)
    if label:
        txt(s, l + 0.26, t + 0.10, w - 0.40, 0.30, label, S_CARDL, True, accent)
        txt(s, l + 0.26, t + 0.44, w - 0.52, h - 0.54, expr, expr_size, False, DARK,
            PP_ALIGN.CENTER, italic=True, anchor=MSO_ANCHOR.MIDDLE)
    else:
        txt(s, l + 0.26, t + 0.06, w - 0.52, h - 0.12, expr, expr_size, False, DARK,
            PP_ALIGN.CENTER, italic=True, anchor=MSO_ANCHOR.MIDDLE)


# ═══════════════════════════════════════════════════════════════════════════════
# LAYER 4 · STRUCTURAL
# Slide-level factories. Every content slide is built via base().
#
# base(active, title, claim)  →  standard content slide
#   active: 0–4 SECTIONS index (highlights that progbar segment)
#   title:  T1 heading (28 pt)
#   claim:  italic sub-heading below gold underline (optional)
#
# section(num, label, title, question, sub)  →  full-bleed section separator
#   num: display string shown large on left panel ("00"–"04")
#   label: UPPERCASE caption below the number
#   title: section name (right-panel heading)
#   question: guiding question displayed on right
#   sub: optional footnote on right (e.g. chapter reference)
#
# take(s, text, y)    →  navy takeaway bar (default y=6.40)
# caption(s, text, y) →  gold-rule caption line
# flatcol(s, ...)     →  label + body column without card border
# add_table(s, ...)   →  styled data table
# ═══════════════════════════════════════════════════════════════════════════════

def progbar(s, active):
    # 5-segment navigation bar at y=0.23. active ∈ {0,1,2,3,4} = SECTIONS index.
    x0, y0, total = 0.55, 0.23, 11.05
    seg = (total - 4*0.08) / 5
    for i, name in enumerate(SECTIONS):
        lx = x0 + i * (seg + 0.08)
        on = (i == active)
        c = chip(s, lx, y0, seg, 0.26, name,
                 NAVY if on else WHITE, WHITE if on else MUTED,
                 size=S_PROG, bold=on, bd=None if on else HAIR)
        soft_shadow(c, alpha=(46 if on else 20))


def footer(s):
    hline(s, 0.55, FOOT_RULE, 12.23, GRAY_BD, 0.8)
    txt(s, 0.55, FOOT_RULE + 0.04, 9.0, 0.26,
        "Towards Scalable and Generalizable Video Retrieval", S_FOOT, False, MUTED)
    txt(s, 10.3, FOOT_RULE + 0.04, 2.48, 0.26, "Zhuoning Guo · HKUST(GZ)", S_FOOT, False, MUTED, PP_ALIGN.RIGHT)


PAGE = {"n": 0}
def _pg(): PAGE["n"] += 1; return PAGE["n"]


def base(active, title, claim=None):
    s = prs.slides.add_slide(BLANK); _bg(s)
    progbar(s, active)
    txt(s, 11.75, 0.22, 1.03, 0.30, str(_pg()), S_FOOT, False, MUTED, PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 0.55, 0.66, 12.23, 0.62, title, S_TITLE, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    hline(s, 0.57, 1.36, 2.10, GOLD, 3.0)
    if claim:
        txt(s, 0.55, 1.42, 12.23, 0.36, claim, S_CLAIM, False, MUTED, italic=True)
    footer(s)
    return s


def take(s, text, y=6.40):
    rect(s, 0.55, y, 12.23, 0.54, NAVY, rounded=True)
    rich(s, 0.80, y + 0.02, 11.8, 0.50, [("Takeaway   ", True, GOLD), (text, False, WHITE)],
         size=S_TAKE, anchor=MSO_ANCHOR.MIDDLE)


def caption(s, text, y=6.52, w=12.23):
    hline(s, 0.57, y, 1.6, GOLD, 1.4)
    txt(s, 0.55, y + 0.09, w, 0.36, text, 12, False, DARK, italic=False)


def flatcol(s, l, t, w, label, body, accent=NAVY, body_size=S_CARDB, bh=1.20):
    # Label + body column without a card border. Use for pipeline-stage descriptions.
    rect(s, l, t + 0.02, 0.08, 0.30, accent)
    txt(s, l + 0.20, t, w - 0.20, 0.34, label, S_CARDL, True, NAVY)
    if isinstance(body, list):
        bullets(s, l + 0.20, t + 0.42, w - 0.24, bh, body, body_size, MUTED, spacing=4)
    else:
        txt(s, l + 0.20, t + 0.42, w - 0.24, bh, body, body_size, False, MUTED)


def add_table(s, headers, rows, l, t, w, h, best_row=-1, hi_col=None,
              col_widths=None, fs=14, hfs=14):
    # best_row: 0-indexed row to highlight GREEN (our method row)
    # hi_col: set of column indices to bold-NAVY (key metric columns)
    # col_widths: list of fractions summing to 1.0
    nc, nr = len(headers), len(rows) + 1
    shp = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(h))
    tb = shp.table
    try: tb.first_row = False; tb.horz_banding = False
    except Exception: pass
    if col_widths:
        tot = Inches(w)
        for j, cwd in enumerate(col_widths): tb.columns[j].width = int(tot * cwd)
    for j, hd in enumerate(headers):
        c = tb.cell(0, j); c.text = hd; c.fill.solid(); c.fill.fore_color.rgb = NAVY
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in c.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT
            p.font.size = Pt(hfs); p.font.bold = True; p.font.color.rgb = WHITE; p.font.name = FONT
    for i, row in enumerate(rows):
        best = (i == best_row)
        for j, val in enumerate(row):
            c = tb.cell(i + 1, j); c.text = str(val); c.fill.solid()
            c.fill.fore_color.rgb = GREEN_L if best else (GRAY_L if i % 2 else WHITE)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in c.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT
                p.font.size = Pt(fs); p.font.name = FONT
                if best: p.font.bold = True; p.font.color.rgb = GREEN
                elif hi_col and j in hi_col: p.font.bold = True; p.font.color.rgb = NAVY
                else: p.font.color.rgb = DARK
    return tb


def section(num, label, title, question, sub=None):
    # Full-bleed section separator. num = display string ("00"–"04").
    s = prs.slides.add_slide(BLANK); _bg(s)
    rect(s, 0, 0, 4.6, 7.50, NAVY2)
    rect(s, 4.6, 0, 0.10, 7.50, GOLD)
    txt(s, 0.45, 1.30, 3.95, 2.0, num, S_SECNUM, True, GOLD, anchor=MSO_ANCHOR.MIDDLE)
    hline(s, 0.58, 3.55, 1.75, WHITE, 2.0)
    txt(s, 0.58, 3.76, 3.7, 0.4, label.upper(), S_FOOT, True, CREAM)
    txt(s, 0.58, 4.18, 3.75, 1.7, title, S_SECTL, True, WHITE)
    txt(s, 5.35, 2.55, 7.4, 0.4, "GUIDING QUESTION", S_FOOT, True, GOLD)
    txt(s, 5.35, 3.02, 7.5, 2.3, question, 20, False, DARK)
    if sub:
        txt(s, 5.35, 5.25, 7.4, 1.4, sub, S_CARDB, False, MUTED)
    txt(s, 11.75, 0.22, 1.03, 0.30, str(_pg()), S_FOOT, False, MUTED, PP_ALIGN.RIGHT)
    notes(s, question)
    return s


# ═══════════════════════════════════════════════════════════════════════════════
# LAYER 5 · APPENDIX
# bslide() is the appendix equivalent of base(): no progbar, "APPENDIX · BACKUP"
# header, B-numbered pages. The Appendix section separator is built inline in
# the deck builder (not a reusable function — it contains deck-specific content).
# ═══════════════════════════════════════════════════════════════════════════════

BP = {"n": 0}
def _bp(): BP["n"] += 1; return "B" + str(BP["n"])

def bslide(title, claim=None):
    s = prs.slides.add_slide(BLANK); _bg(s)
    txt(s, 0.55, 0.30, 10.0, 0.34, "APPENDIX · BACKUP", S_FOOT, True, GOLD2)
    txt(s, 11.75, 0.30, 1.03, 0.30, _bp(), S_FOOT, False, MUTED, PP_ALIGN.RIGHT)
    txt(s, 0.55, 0.66, 12.23, 0.62, title, S_TITLE, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    hline(s, 0.57, 1.36, 2.10, GOLD, 3.0)
    if claim: txt(s, 0.55, 1.42, 12.23, 0.36, claim, S_CLAIM, False, MUTED, italic=True)
    footer(s)
    return s


# ═══════════════════════════════════════════════════════════════════════════════
# LAYER 6 · UTILITIES
# _force/_tf/_walk: walk every shape and hard-pin font to FONT (Times New Roman).
#   Call `for slide in prs.slides: _walk(slide.shapes)` after all slides are built.
#
# _register_notes_master: fixes PowerPoint's "found a problem / repair?" dialog.
#   Root cause: python-pptx writes the notesMaster relationship into .rels but
#   omits the <p:notesMasterIdLst> element in presentation.xml, leaving a dangling
#   reference. This function adds the missing element to make the package consistent.
#   Call once: _register_notes_master(prs) before prs.save().
# ═══════════════════════════════════════════════════════════════════════════════

def _force(rpr, name=FONT):
    for tag in ("a:latin", "a:ea", "a:cs"):
        for e in rpr.findall(qn(tag)): rpr.remove(e)
    for tag in ("a:latin", "a:ea", "a:cs"):
        rpr.append(rpr.makeelement(qn(tag), {"typeface": name}))

def _tf(tf):
    for p in tf.paragraphs:
        pPr = p._p.get_or_add_pPr()
        d = pPr.find(qn("a:defRPr"))
        if d is None:
            d = pPr.makeelement(qn("a:defRPr"), {}); pPr.append(d)
        _force(d)
        for r in p.runs: _force(r._r.get_or_add_rPr())

def _walk(shapes):
    for sh in shapes:
        if sh.has_text_frame: _tf(sh.text_frame)
        if sh.shape_type == 6: _walk(sh.shapes)
        try:
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells: _tf(c.text_frame)
        except Exception: pass

def delete_slide(prs, idx):
    # python-pptx has no public delete API; reach into internals.
    sldIdLst = prs.slides._sldIdLst
    ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rId = sldIdLst[idx].attrib[f"{{{ns}}}id"]
    del prs.part._rels._rels[rId]
    sldIdLst.remove(sldIdLst[idx])


def _register_notes_master(prs):
    RT_NM = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesMaster"
    nm_rid = next((r.rId for r in prs.part.rels.values() if r.reltype == RT_NM), None)
    if not nm_rid:
        return False
    sldIdLst = prs.slides._sldIdLst
    pres_el = sldIdLst.getparent()
    if pres_el.find(qn("p:notesMasterIdLst")) is not None:
        return False
    nmil = pres_el.makeelement(qn("p:notesMasterIdLst"), {})
    nmid = nmil.makeelement(qn("p:notesMasterId"), {qn("r:id"): nm_rid})
    nmil.append(nmid)
    sldIdLst.addprevious(nmil)
    return True


# ── Smoke test ────────────────────────────────────────────────────────────────
# Running this file directly produces an empty PPTX (no slides, design system
# only) to verify all imports and function definitions load without errors.
if __name__ == "__main__":
    for slide in prs.slides: _walk(slide.shapes)
    _register_notes_master(prs)
    prs.save("design-system-smoke.pptx")
    print(f"Smoke test OK — slides: {len(prs.slides)}")
